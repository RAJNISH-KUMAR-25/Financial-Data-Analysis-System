import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, LineChart, Reference
from datetime import datetime
import random

random.seed(42)

def style_header(ws, row, num_cols, color="1B3A6B"):
    fill = PatternFill("solid", fgColor=color)
    for col in range(1, num_cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = fill
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center")

def style_row(ws, row, num_cols, color="D6E4F0"):
    fill = PatternFill("solid", fgColor=color)
    for col in range(1, num_cols + 1):
        ws.cell(row=row, column=col).fill = fill

def set_col_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

def build_report():
    wb = Workbook()

    # ── Sheet 1: Monthly P&L ──
    ws1 = wb.active
    ws1.title = "Monthly P&L"
    ws1["A1"] = "Financial Data Analysis System - Monthly P&L Report"
    ws1["A1"].font = Font(bold=True, size=16, color="1B3A6B")
    ws1["A2"] = f"Generated: {datetime.now().strftime('%Y-%m-%d')}"
    ws1.append([])

    months = ["Jan", "Feb", "Mar", "Apr", "May", "Jun",
              "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
    headers = ["Month", "Revenue ($)", "Expense ($)", "Net Profit ($)", "Margin (%)"]
    ws1.append(headers)
    style_header(ws1, 4, 5)

    for i, m in enumerate(months):
        rev = round(random.uniform(200000, 500000), 2)
        exp = round(random.uniform(150000, 350000), 2)
        net = round(rev - exp, 2)
        margin = round((net / rev) * 100, 2)
        ws1.append([m, rev, exp, net, margin])
        if i % 2 == 0:
            style_row(ws1, 5 + i, 5, "EAF2FB")

    set_col_widths(ws1, [14, 18, 18, 18, 14])

    # Line chart
    chart = LineChart()
    chart.title = "Monthly Revenue vs Expense"
    chart.style = 10
    data = Reference(ws1, min_col=2, max_col=3, min_row=4, max_row=16)
    chart.add_data(data, titles_from_data=True)
    cats = Reference(ws1, min_col=1, min_row=5, max_row=16)
    chart.set_categories(cats)
    ws1.add_chart(chart, "G4")

    # ── Sheet 2: Department Budget ──
    ws2 = wb.create_sheet("Department Budget")
    depts = ["Sales", "Marketing", "Operations", "HR", "IT", "Finance"]
    headers2 = ["Department", "Budget ($)", "Actual ($)", "Variance ($)", "Utilization (%)"]
    ws2.append(headers2)
    style_header(ws2, 1, 5)

    for i, d in enumerate(depts):
        budget = round(random.uniform(100000, 500000), 2)
        actual = round(budget * random.uniform(0.7, 1.3), 2)
        var    = round(actual - budget, 2)
        util   = round((actual / budget) * 100, 2)
        ws2.append([d, budget, actual, var, util])
        if i % 2 == 0:
            style_row(ws2, 2 + i, 5, "EAF2FB")

    set_col_widths(ws2, [18, 18, 18, 18, 18])

    # ── Sheet 3: Forecast ──
    ws3 = wb.create_sheet("Revenue Forecast")
    headers3 = ["Period", "Forecasted Revenue ($)", "Lower Bound ($)", "Upper Bound ($)", "Confidence (%)"]
    ws3.append(headers3)
    style_header(ws3, 1, 5)

    base = 350000
    periods = ["Jan 2025", "Feb 2025", "Mar 2025", "Apr 2025", "May 2025", "Jun 2025"]
    for i, p in enumerate(periods):
        forecast = round(base * (1 + 0.03 * (i + 1)) + random.uniform(-5000, 5000), 2)
        lower    = round(forecast * 0.92, 2)
        upper    = round(forecast * 1.08, 2)
        conf     = round(random.uniform(85, 95), 1)
        ws3.append([p, forecast, lower, upper, conf])
        if i % 2 == 0:
            style_row(ws3, 2 + i, 5, "EAF2FB")

    set_col_widths(ws3, [16, 24, 20, 20, 16])

    # ── Sheet 4: Statistical Summary ──
    ws4 = wb.create_sheet("Statistical Summary")
    headers4 = ["Category", "Count", "Mean ($)", "Median ($)", "Std Dev ($)", "Min ($)", "Max ($)", "Total ($)"]
    ws4.append(headers4)
    style_header(ws4, 1, 8)

    categories = ["Revenue", "Expense", "Investment", "Tax", "Refund"]
    for i, cat in enumerate(categories):
        count  = random.randint(8000, 20000)
        mean   = round(random.uniform(1000, 8000), 2)
        median = round(mean * random.uniform(0.6, 0.9), 2)
        std    = round(mean * random.uniform(0.3, 0.8), 2)
        mn     = round(mean * 0.05, 2)
        mx     = round(mean * random.uniform(5, 15), 2)
        total  = round(count * mean, 2)
        ws4.append([cat, count, mean, median, std, mn, mx, total])
        if i % 2 == 0:
            style_row(ws4, 2 + i, 8, "EAF2FB")

    set_col_widths(ws4, [16, 10, 14, 14, 14, 12, 14, 18])

    wb.save("reports/Financial_KPI_Report.xlsx")
    print("Excel report saved → reports/Financial_KPI_Report.xlsx")

if __name__ == "__main__":
    build_report()
